"""What Claude sees of a file to edit it, and how its changes are applied.

text (txt, md): numbered lines; exact find-and-replace edits or new content.
  The byte-order mark and Windows/Unix line endings are kept as they were.
docx: body paragraphs (p1, p2, ...) and tables (t1, cells t1.r2.c3) in
  order. python-docx edits the document's XML in place, so headers, footers,
  images and everything a change doesn't touch stay exactly as they were;
  replacing text works on the runs, so bold, italics and links around it
  survive. Ids are those of the version that was opened: within one call every
  operation refers to that numbering, inserts and deletes included.
xlsx: sheets and cells (Rates!B4) with values and formulas. openpyxl drops
  charts, pivot tables, pictures and some other parts when it saves, so a
  workbook containing any of them is refused instead of being damaged.
  Inserting or deleting rows doesn't update formulas that point past them, so
  it's refused in workbooks with formulas; renaming a sheet other formulas
  name is refused too. Excel recalculates every formula when it next opens.
"""

from __future__ import annotations

import copy
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass

from mnemerynne.editing import EditError

MAX_TEXT_LINES = 1000
MAX_DOC_BLOCKS = 300
MAX_SHEET_CELLS = 2000


@dataclass
class Result:
    data: bytes
    changes: list[str]


def _need(op: dict, *keys: str):
    missing = [k for k in keys if op.get(k) in (None, "")]
    if missing:
        raise EditError(f"Operation {op.get('op')!r} needs {', '.join(missing)}.")
    return [op[k] for k in keys]


def _count(ops, limit: int) -> list[dict]:
    if not ops:
        raise EditError("Give at least one operation.")
    if len(ops) > limit:
        raise EditError(f"At most {limit} operations per call; split the change into several calls.")
    if not all(isinstance(o, dict) for o in ops):
        raise EditError("Each operation is an object with an `op` field.")
    return list(ops)


# --------------------------------------------------------------------- text
def _decode(data: bytes) -> tuple[str, bool, str]:
    """(text with \\n line ends, had a BOM, the line ending to write back)."""
    bom = data.startswith(b"\xef\xbb\xbf")
    try:
        text = data.decode("utf-8-sig" if bom else "utf-8")
    except UnicodeDecodeError:
        raise EditError("This file isn't UTF-8 text, so it can't be edited safely here.") from None
    crlf = text.count("\r\n")
    newline = "\r\n" if crlf and crlf >= text.count("\n") - crlf else "\n"
    return text.replace("\r\n", "\n"), bom, newline


def _encode(text: str, bom: bool, newline: str) -> bytes:
    if newline != "\n":
        text = text.replace("\n", newline)
    return (b"\xef\xbb\xbf" if bom else b"") + text.encode("utf-8")


def text_view(data: bytes, start_line: int = 1, max_lines: int = 400) -> dict:
    text, _, newline = _decode(data)
    lines = text.split("\n")
    start = max(1, int(start_line))
    count = max(1, min(int(max_lines), MAX_TEXT_LINES))
    shown = lines[start - 1:start - 1 + count]
    width = len(str(start + len(shown)))
    return {"format": "text", "total_lines": len(lines), "first_line": start, "last_line": start + len(shown) - 1,
            "line_endings": "windows" if newline == "\r\n" else "unix",
            "text": "\n".join(f"{n:>{width}}| {line}" for n, line in enumerate(shown, start)),
            "note": "Line numbers are for reference only; edits match the text itself."}


def text_apply(data: bytes, edits: list[dict] | None = None, content: str | None = None,
               max_ops: int = 50) -> Result:
    text, bom, newline = _decode(data)
    if content is not None:
        if edits:
            raise EditError("Give either `content` (the whole new text) or `edits`, not both.")
        return Result(_encode(str(content).replace("\r\n", "\n"), bom, newline), ["replaced the whole text"])
    changes = []
    for n, edit in enumerate(_count(edits, max_ops), 1):
        old, new = edit.get("old"), edit.get("new")
        if not isinstance(old, str) or not old or not isinstance(new, str):
            raise EditError(f"Edit {n} needs `old` (the exact text to replace) and `new`.")
        old, new = old.replace("\r\n", "\n"), new.replace("\r\n", "\n")
        found = text.count(old)
        if found != 1:
            raise EditError(f"Edit {n}: the `old` text appears {found} times; it must appear exactly once. "
                            + ("Copy it exactly from open_file." if not found else "Include more of the surrounding text."))
        text = text.replace(old, new, 1)
        changes.append(f"replaced {_short(old)!r} with {_short(new)!r}")
    return Result(_encode(text, bom, newline), changes)


def text_new(content: str) -> bytes:
    return str(content).replace("\r\n", "\n").encode("utf-8")


def _short(s: str, n: int = 40) -> str:
    s = " ".join(s.split())
    return s if len(s) <= n else s[: n - 1] + "…"


# --------------------------------------------------------------------- docx
def _docx():
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    return docx, Paragraph, Table


def _blocks(document):
    """Body paragraphs and tables in order, with their ids."""
    _, Paragraph, Table = _docx()
    out, p_n, t_n = [], 0, 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p_n += 1
            out.append((f"p{p_n}", Paragraph(child, document._body)))
        elif tag == "tbl":
            t_n += 1
            out.append((f"t{t_n}", Table(child, document._body)))
    return out


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def docx_view(data: bytes, start: int = 1, max_blocks: int = 200) -> dict:
    docx, Paragraph, _ = _docx()
    document = docx.Document(io.BytesIO(data))
    blocks = _blocks(document)
    first = max(1, int(start))
    count = max(1, min(int(max_blocks), MAX_DOC_BLOCKS))
    shown = []
    for block_id, block in blocks[first - 1:first - 1 + count]:
        if isinstance(block, Paragraph):
            shown.append({"id": block_id, "style": block.style.name if block.style is not None else None,
                          "text": block.text})
        else:
            shown.append({"id": block_id, "table": [[{"id": f"{block_id}.r{r}.c{c}", "text": _cell_text(cell)}
                                                     for c, cell in enumerate(row.cells, 1)]
                                                    for r, row in enumerate(block.rows, 1)]})
    used = sorted({b.style.name for _, b in blocks if isinstance(b, Paragraph) and b.style is not None})
    return {"format": "docx", "total_blocks": len(blocks), "first_block": first,
            "last_block": first + len(shown) - 1, "blocks": shown, "styles_in_use": used,
            "note": "Paragraph ids (p…) count body paragraphs only; table ids (t…) count tables."}


def _style(document, name):
    if not name:
        return None
    try:
        return document.styles[name]
    except KeyError:
        raise EditError(f"There's no style called {name!r} in this document.") from None


def _runs_text(paragraph) -> str:
    return "".join(r.text for r in paragraph.runs)


def _replace_in_runs(paragraph, old: str, new: str) -> None:
    """Replace `old` (exactly once in the paragraph's runs) keeping formatting:
    the new text takes the formatting of the run where the old text began."""
    runs = paragraph.runs
    joined = _runs_text(paragraph)
    start = joined.find(old)
    end = start + len(old)
    pos = 0
    first = True
    for run in runs:
        text = run.text
        r_start, r_end = pos, pos + len(text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        keep_before = text[: max(0, start - r_start)] if first else ""
        keep_after = text[end - r_start:] if r_end > end else ""
        run.text = keep_before + (new if first else "") + keep_after
        first = False


def _clear_paragraph(paragraph):
    """Remove the paragraph's content (runs, links, fields) but keep its
    properties; returns the first run's formatting to reuse."""
    p = paragraph._p
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_rpr = copy.deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(p):
        if child.tag.rsplit("}", 1)[-1] != "pPr":
            p.remove(child)
    return first_rpr


def _set_text(paragraph, text: str) -> None:
    rpr = _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def _set_cell(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _set_text(paragraphs[0], text)


def docx_apply(data: bytes, ops: list[dict], max_ops: int = 50) -> Result:
    docx, Paragraph, Table = _docx()
    from docx.oxml import OxmlElement

    ops = _count(ops, max_ops)
    document = docx.Document(io.BytesIO(data))
    blocks = dict(_blocks(document))
    deleted: set[str] = set()

    def block(block_id: str, want=None):
        if block_id in deleted:
            raise EditError(f"{block_id} was deleted earlier in this call.")
        found = blocks.get(str(block_id))
        if found is None or (want is not None and not isinstance(found, want)):
            kind = "paragraph" if want is Paragraph else "table" if want is Table else "block"
            raise EditError(f"There's no {kind} {block_id!r} in the version that was opened.")
        return found

    def cell(ref: str):
        m = re.fullmatch(r"(t\d+)\.r(\d+)\.c(\d+)", str(ref))
        if not m:
            raise EditError(f"Cells are named like t1.r2.c3, not {ref!r}.")
        table = block(m.group(1), Table)
        r, c = int(m.group(2)), int(m.group(3))
        if not (1 <= r <= len(table.rows) and 1 <= c <= len(table.columns)):
            raise EditError(f"{ref} is outside {m.group(1)} ({len(table.rows)} rows × {len(table.columns)} columns).")
        return table.cell(r - 1, c - 1)

    changes = []
    for n, op in enumerate(ops, 1):
        kind = op.get("op")
        if kind == "replace_text":
            pid, old = _need(op, "paragraph", "old")
            new = str(op.get("new") or "")
            p = block(pid, Paragraph)
            found = _runs_text(p).count(old)
            if found != 1:
                if found == 0 and old in p.text:
                    raise EditError(f"Operation {n}: in {pid} that text is inside a link or field; "
                                    f"use set_paragraph to rewrite the paragraph instead.")
                raise EditError(f"Operation {n}: {old!r} appears {found} times in {pid}; it must appear once.")
            _replace_in_runs(p, old, new)
            changes.append(f"{pid}: replaced {_short(old)!r} with {_short(new)!r}")
        elif kind == "set_paragraph":
            (pid,) = _need(op, "paragraph")
            p = block(pid, Paragraph)
            _set_text(p, str(op.get("text") or ""))
            if op.get("style"):
                p.style = _style(document, op["style"])
            changes.append(f"{pid}: rewrote the paragraph")
        elif kind == "insert_paragraph":
            text = str(op.get("text") or "")
            after, before = op.get("after"), op.get("before")
            if bool(after) == bool(before):
                raise EditError(f"Operation {n}: give exactly one of `after` or `before` (a p…/t… id, or \"end\").")
            anchor_id = after or before
            if anchor_id == "end":
                new_p = document.add_paragraph(text)
            else:
                anchor = block(anchor_id)
                element = OxmlElement("w:p")
                (anchor._element.addnext if after else anchor._element.addprevious)(element)
                new_p = Paragraph(element, document._body)
                if isinstance(anchor, Paragraph) and anchor._p.pPr is not None:
                    element.insert(0, copy.deepcopy(anchor._p.pPr))   # same style and spacing as its neighbour
                new_p.add_run(text)
            if op.get("style"):
                new_p.style = _style(document, op["style"])
            changes.append(f"inserted a paragraph {'after' if after else 'before'} {anchor_id}: {_short(text)!r}")
        elif kind == "delete_paragraph":
            (pid,) = _need(op, "paragraph")
            p = block(pid, Paragraph)
            p._p.getparent().remove(p._p)
            deleted.add(pid)
            changes.append(f"deleted {pid}: {_short(p.text)!r}")
        elif kind == "set_cell":
            ref, = _need(op, "cell")
            _set_cell(cell(ref), str(op.get("text") or ""))
            changes.append(f"{ref}: set to {_short(str(op.get('text') or ''))!r}")
        elif kind == "add_table_row":
            (tid,) = _need(op, "table")
            table = block(tid, Table)
            values = op.get("cells") or []
            if not isinstance(values, list):
                raise EditError(f"Operation {n}: `cells` is a list of the new row's texts.")
            after_row = int(op.get("after_row") or len(table.rows))
            if not 1 <= after_row <= len(table.rows):
                raise EditError(f"Operation {n}: {tid} has {len(table.rows)} rows.")
            source = table.rows[after_row - 1]._tr
            new_tr = copy.deepcopy(source)   # same cell widths, borders and formatting
            source.addnext(new_tr)
            new_row = next(r for r in table.rows if r._tr is new_tr)
            for c, target in enumerate(new_row.cells):
                _set_cell(target, str(values[c]) if c < len(values) else "")
            changes.append(f"{tid}: added a row after row {after_row}")
        else:
            raise EditError(f"Operation {n}: unknown op {kind!r}. Use replace_text, set_paragraph, insert_paragraph, "
                            f"delete_paragraph, set_cell or add_table_row.")
    out = io.BytesIO()
    document.save(out)
    return Result(out.getvalue(), changes)


def docx_new(paragraphs: list) -> bytes:
    docx, _, _ = _docx()
    document = docx.Document()
    for item in paragraphs or [""]:
        if isinstance(item, dict):
            p = document.add_paragraph(str(item.get("text") or ""))
            if item.get("style"):
                p.style = _style(document, item["style"])
        else:
            document.add_paragraph(str(item))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --------------------------------------------------------------------- xlsx
# Parts openpyxl would lose or damage when it saves.
_UNSAFE_PARTS = [
    (re.compile(r"xl/charts/"), "charts"),
    (re.compile(r"xl/pivot(Tables|Cache)/"), "pivot tables"),
    (re.compile(r"xl/drawings/drawing\d+\.xml$"), "pictures, shapes or charts"),
    (re.compile(r"xl/media/"), "pictures"),
    (re.compile(r"xl/slicers?|xl/slicerCaches/"), "slicers"),
    (re.compile(r"xl/vbaProject"), "macros"),
    (re.compile(r"xl/threadedComments/"), "threaded comments"),
    (re.compile(r"xl/ctrlProps/|xl/activeX/"), "form controls"),
    (re.compile(r"xl/richData/"), "pictures in cells"),
]


def _unsafe(data: bytes) -> list[str]:
    try:
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile:
        raise EditError("This doesn't look like a valid .xlsx file.") from None
    return sorted({label for pattern, label in _UNSAFE_PARTS for n in names if pattern.match(n)})


def _workbook(data: bytes, values: bool = False):
    import openpyxl
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=values)
    except Exception as exc:   # openpyxl raises many kinds for broken files
        raise EditError(f"Couldn't read the workbook: {exc}") from None


def _plain(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def _sheet(wb, name):
    if name is None:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise EditError(f"There's no sheet {name!r}; the sheets are {', '.join(wb.sheetnames)}.")
    return wb[name]


def xlsx_view(data: bytes, sheet: str | None = None, cells: str | None = None) -> dict:
    from openpyxl.utils import get_column_letter, range_boundaries

    wb, cached = _workbook(data), _workbook(data, values=True)
    ws = _sheet(wb, sheet)
    if cells:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cells.split("!")[-1])
        except ValueError:
            raise EditError(f"Name cells like A1:F40, not {cells!r}.") from None
    else:
        min_col, min_row = 1, 1
        max_col = min(ws.max_column, 26)
        max_row = min(ws.max_row, max(1, MAX_SHEET_CELLS // max(max_col, 1)))
    if (max_col - min_col + 1) * (max_row - min_row + 1) > MAX_SHEET_CELLS:
        raise EditError(f"Ask for at most {MAX_SHEET_CELLS} cells at a time (e.g. A1:J200).")
    rows = []
    for r in range(min_row, max_row + 1):
        row = {}
        for c in range(min_col, max_col + 1):
            ref = f"{get_column_letter(c)}{r}"
            value = ws[ref].value
            if value is None:
                continue
            if isinstance(value, str) and value.startswith("="):
                row[ref] = {"formula": value, "value": _plain(cached[ws.title][ref].value)}
            else:
                row[ref] = _plain(value)
        if row:
            rows.append(row)
    unsafe = _unsafe(data)
    return {"format": "xlsx", "sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column}
                                         for s in wb.worksheets],
            "sheet": ws.title, "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
            "cells": rows,
            "editable": not unsafe,
            "note": (f"Can't be edited here: it contains {', '.join(unsafe)}, which saving would damage."
                     if unsafe else "Formulas show with their last calculated value; Excel recalculates on open.")}


def _has_formulas(wb) -> bool:
    return any(isinstance(c.value, str) and c.value.startswith("=")
               for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def _cell_value(value):
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    return str(value)


def xlsx_apply(data: bytes, ops: list[dict], max_ops: int = 50, max_cells: int = 200_000) -> Result:
    unsafe = _unsafe(data)
    if unsafe:
        raise EditError(f"This workbook contains {', '.join(unsafe)}; saving it here would lose them, "
                        f"so it can't be edited by Claude. Make the change in Excel.")
    ops = _count(ops, max_ops)
    wb = _workbook(data)
    if sum(ws.max_row * ws.max_column for ws in wb.worksheets) > max_cells:
        raise EditError(f"This workbook is too large to edit here (over {max_cells:,} cells).")
    changes = []
    for n, op in enumerate(ops, 1):
        kind = op.get("op")
        if kind == "set_cells":
            values = op.get("cells")
            if not isinstance(values, dict) or not values:
                raise EditError(f"Operation {n}: `cells` maps cell names to values, e.g. {{\"B4\": 12.5}}.")
            ws = _sheet(wb, op.get("sheet"))
            for ref, value in values.items():
                if not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d{0,6}", str(ref)):
                    raise EditError(f"Operation {n}: {ref!r} isn't a cell name like B4.")
                ws[str(ref).upper()] = _cell_value(value)
            changes.append(f"{ws.title}: set {', '.join(list(values)[:8])}{'…' if len(values) > 8 else ''}")
        elif kind == "append_rows":
            rows = op.get("rows")
            if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
                raise EditError(f"Operation {n}: `rows` is a list of rows, each a list of values.")
            ws = _sheet(wb, op.get("sheet"))
            for row in rows:
                ws.append([_cell_value(v) for v in row])
            changes.append(f"{ws.title}: added {len(rows)} row(s) at the end")
        elif kind in ("insert_rows", "delete_rows"):
            ws = _sheet(wb, op.get("sheet"))
            if _has_formulas(wb):
                raise EditError(f"Operation {n}: this workbook has formulas, and {kind} here wouldn't update the "
                                f"ones that point past the change. Use append_rows or set_cells, or do it in Excel.")
            (row,) = _need(op, "row")
            count = max(1, int(op.get("count") or 1))
            (ws.insert_rows if kind == "insert_rows" else ws.delete_rows)(int(row), count)
            changes.append(f"{ws.title}: {'inserted' if kind == 'insert_rows' else 'deleted'} {count} row(s) at {row}")
        elif kind == "add_sheet":
            (name,) = _need(op, "name")
            if name in wb.sheetnames:
                raise EditError(f"Operation {n}: there's already a sheet called {name!r}.")
            wb.create_sheet(str(name))
            changes.append(f"added sheet {name!r}")
        elif kind == "rename_sheet":
            old, new = _need(op, "sheet", "name")
            ws = _sheet(wb, old)
            mentions = re.compile(rf"(?:'{re.escape(old)}'|\b{re.escape(old)})!")
            if any(isinstance(c.value, str) and c.value.startswith("=") and mentions.search(c.value)
                   for s in wb.worksheets for r in s.iter_rows() for c in r):
                raise EditError(f"Operation {n}: formulas refer to {old!r}, and renaming here wouldn't update them.")
            ws.title = str(new)
            changes.append(f"renamed sheet {old!r} to {new!r}")
        else:
            raise EditError(f"Operation {n}: unknown op {kind!r}. Use set_cells, append_rows, insert_rows, "
                            f"delete_rows, add_sheet or rename_sheet.")
    wb.calculation.fullCalcOnLoad = True   # cached results of changed formulas are stale until Excel recalculates
    out = io.BytesIO()
    wb.save(out)
    return Result(out.getvalue(), changes)


def xlsx_new(rows: list, sheet: str | None = None) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if sheet:
        ws.title = str(sheet)
    for row in rows or []:
        if not isinstance(row, list):
            raise EditError("`rows` is a list of rows, each a list of values.")
        ws.append([_cell_value(v) for v in row])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
