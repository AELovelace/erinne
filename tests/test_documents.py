import datetime
import io
import zipfile

import pytest
from fastapi.testclient import TestClient

from fakes import FakeEmbedder
from mnemerynne import users, worker
from mnemerynne.ingest import documents, store
from mnemerynne.web.app import Services, create_app


def _paths(segments):
    return [s.section_path for s in segments]


def _docx(path):
    import docx

    d = docx.Document()
    d.add_paragraph("Front desk handbook")
    d.add_heading("Night audit", level=1)
    d.add_paragraph("Run the night audit at 2am.")
    d.add_heading("Steps", level=2)
    d.add_paragraph("Close the shift, then roll the business date.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Report", "Printer"
    table.cell(1, 0).text, table.cell(1, 1).text = "Trial balance", "Back office"
    d.add_heading("Checkout", level=1)
    d.add_paragraph("Late checkout until 1pm for members.")
    d.save(path)
    return path


def test_word_documents_keep_their_headings(tmp_path):
    segs = documents.segments(_docx(tmp_path / "h.docx"), "h.docx")
    assert _paths(segs) == ["", "Night audit", "Night audit > Steps", "Checkout"]
    assert "Trial balance | Back office" in segs[2].text
    assert [s.anchor for s in segs] == ["s1", "s2", "s3", "s4"]


def test_presentations_are_read_slide_by_slide(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Welcome"
    slide.placeholders[1].text = "Greet every guest by name."
    slide.notes_slide.notes_text_frame.text = "Smile."
    deck.slides.add_slide(deck.slide_layouts[6])    # blank: skipped
    path = tmp_path / "t.pptx"
    deck.save(path)
    segs = documents.segments(path, "t.pptx")
    assert _paths(segs) == ["Slide 1: Welcome"]
    assert "Greet every guest by name." in segs[0].text and "Notes: Smile." in segs[0].text


def test_spreadsheets_are_read_in_row_blocks_with_column_names(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Rates"
    sheet.append(["Room", "Rate", "From"])
    for n in range(1, 41):
        sheet.append([f"Room {n}", 100 + n, datetime.date(2026, 11, 1)])
    path = tmp_path / "r.xlsx"
    book.save(path)
    segs = documents.segments(path, "r.xlsx")
    assert _paths(segs) == ["Rates, rows 2-31", "Rates, rows 32-41"]
    assert segs[0].text.splitlines()[0] == "Room: Room 1 | Rate: 101 | From: 2026-11-01"


def test_csv_files(tmp_path):
    path = tmp_path / "contacts.csv"
    path.write_text("Name;Phone\nMaintenance;x204\nSecurity;x911\n", encoding="utf-8")
    segs = documents.segments(path, "contacts.csv")
    assert _paths(segs) == ["contacts, rows 2-3"] and "Name: Security | Phone: x911" in segs[0].text


def _odf(path, body):
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content '
               'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
               'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0">'
               f'<office:body>{body}</office:body></office:document-content>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("content.xml", content)
    return path


def test_opendocument_text_presentation_and_sheet(tmp_path):
    odt = _odf(tmp_path / "a.odt", '<office:text><text:h text:outline-level="1">Pool</text:h>'
                                   '<text:p>Open<text:s/>7am to 10pm.</text:p></office:text>')
    assert [(s.section_path, s.text) for s in documents.segments(odt, "a.odt")] == [("Pool", "Open 7am to 10pm.")]
    odp = _odf(tmp_path / "b.odp", '<office:presentation><draw:page draw:name="page1">'
                                   '<text:p>Safety</text:p><text:p>Know the exits.</text:p></draw:page>'
                                   '</office:presentation>')
    assert _paths(documents.segments(odp, "b.odp")) == ["Slide 1: Safety"]
    ods = _odf(tmp_path / "c.ods", '<office:spreadsheet><table:table table:name="Staff">'
                                   '<table:table-row><table:table-cell><text:p>Name</text:p></table:table-cell>'
                                   '</table:table-row><table:table-row><table:table-cell><text:p>Ana</text:p>'
                                   '</table:table-cell></table:table-row>'
                                   '<table:table-row table:number-rows-repeated="1048000"><table:table-cell/>'
                                   '</table:table-row></table:table></office:spreadsheet>')
    segs = documents.segments(ods, "c.ods")
    assert _paths(segs) == ["Staff, rows 2-2"] and segs[0].text == "Name: Ana"


def test_rtf_and_html(tmp_path):
    rtf = tmp_path / "a.rtf"
    rtf.write_text(r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0 Parking is \b free\b0  for guests.\par}")
    assert "Parking is free for guests." in documents.segments(rtf, "a.rtf")[0].text
    html = tmp_path / "a.html"
    html.write_text("<html><head><title>x</title><script>alert(1)</script></head><body>"
                    "<h1>Pets</h1><p>Dogs under 20 kg are welcome.</p><h2>Fees</h2><p>$25 per night.</p></body></html>")
    segs = documents.segments(html, "a.html")
    assert [(s.section_path, s.text) for s in segs] == [
        ("Pets", "Dogs under 20 kg are welcome."), ("Pets > Fees", "$25 per night.")]


def test_code_is_cited_by_lines_and_definition():
    text = "import os\n\n\ndef sync(conn):\n" + "    pass\n" * 70 + "class Stats:\n    x = 1\n"
    segs = documents.code_segments(text)
    assert _paths(segs) == ["lines 1-60 (def sync)", "lines 61-77 (class Stats)"]
    assert documents.kind_of("deploy.ps1") == "code" and documents.kind_of("Dockerfile") == "code"
    assert documents.kind_of("photo.jpg") is None


def test_zip_bombs_and_junk_are_refused(tmp_path):
    junk = tmp_path / "x.docx"
    junk.write_bytes(b"not a zip")
    with pytest.raises(documents.DocumentError, match="damaged"):
        documents.segments(junk, "x.docx")


# ------------------------------------------------------------------ end to end
@pytest.fixture
def client(settings, conn):
    with TestClient(create_app(settings, Services(llm=None, embedder=FakeEmbedder()))) as c:
        yield c


def _login(client, conn):
    users.create_user(conn, "alice", "password123")
    client.post("/login", data={"username": "alice", "password": "password123"})


def test_uploaded_documents_become_searchable_with_a_text_view(client, conn, settings, tmp_path):
    _login(client, conn)
    data = _docx(tmp_path / "handbook.docx").read_bytes()
    up = client.post("/api/files", files={"file": ("handbook.docx", io.BytesIO(data), "application/octet-stream")})
    assert up.status_code == 201 and up.json()["kind"] == "doc" and up.json()["preview"] == "md"
    while worker.run_one(conn, settings, FakeEmbedder()):
        pass
    assert client.get(f"/api/files/{up.json()['id']}").json()["status"] == "ready"
    chunks = conn.execute("SELECT section_path, anchor FROM chunks WHERE file_id = ? ORDER BY ord",
                          (up.json()["id"],)).fetchall()
    assert ("Night audit > Steps", "s3") in [tuple(c) for c in chunks]
    view = client.get(f"/files/{up.json()['id']}/view").text
    assert '<h3 id="s3">Night audit &gt; Steps</h3>' in view and "Run the night audit at 2am." in view

    code = client.post("/api/files", files={"file": ("sync.py", io.BytesIO(b"def sync():\n    return 1\n"), "text/x-python")})
    assert code.json()["kind"] == "code"
    while worker.run_one(conn, settings, FakeEmbedder()):
        pass
    assert "<pre>def sync():" in client.get(f"/files/{code.json()['id']}/view").text


def test_old_word_files_need_libreoffice(client, conn, settings):
    settings.libreoffice_command = "definitely-not-installed"
    _login(client, conn)
    up = client.post("/api/files", files={"file": ("old.doc", io.BytesIO(b"\xd0\xcf\x11\xe0 old"), "application/msword")})
    assert up.json()["status"] == "stored" and "LibreOffice" in up.json()["notes"]


def test_reclassify_picks_up_files_stored_before_support(conn, settings, tmp_path):
    uid = users.create_user(conn, "alice", "password123")
    path = _docx(tmp_path / "old.docx")
    received = store.receive(io.BytesIO(path.read_bytes()), "old.docx", settings)
    received.kind, received.mime = "other", "application/octet-stream"   # as an older version stored it
    from mnemerynne import permissions
    file_id, _ = store.add_file(conn, settings, received, permissions.private_collection_id(conn, uid), uid)
    assert conn.execute("SELECT status FROM files WHERE id = ?", (file_id,)).fetchone()[0] == "stored"
    assert store.reclassify(conn, settings) == ["old.docx"]
    assert conn.execute("SELECT b.kind, f.status FROM files f JOIN blobs b ON b.sha256 = f.blob_sha256 "
                        "WHERE f.id = ?", (file_id,)).fetchone()[:] == ("doc", "queued")
    assert store.reclassify(conn, settings) == []
