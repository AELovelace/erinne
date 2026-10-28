"""Claude editing files: the Agents folder, permissions, versions, formats."""

import io
import json
from urllib.parse import unquote

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from mcp.server.mcpserver.exceptions import ToolError

from fakes import FakeDirectory, FakeEmbedder, FakeShareChecker
from mnemerynne import accounts, api_tokens, db, editing, editing_formats, fs, jobs, permissions, share, users
from mnemerynne.config import Settings
from mnemerynne.ingest import store
from mnemerynne.mcp_server import Tools
from mnemerynne.web.app import Services, create_app


@pytest.fixture
def world(conn, settings):
    alice = users.create_user(conn, "alice", "password123")
    bob = users.create_user(conn, "bob", "password123")
    brand = permissions.create_library(conn, "Brand")
    permissions.grant_user(conn, brand, alice, can_write=True)
    permissions.grant_user(conn, brand, bob)
    return {"alice": alice, "bob": bob, "brand": brand, "inbox": permissions.private_collection_id(conn, alice)}


def _caller(conn, user_id, can_edit=True):
    return api_tokens.verify(conn, api_tokens.create(conn, user_id, "t", can_edit=can_edit)[1])


@pytest.fixture
def tools(settings):
    return Tools(settings, lambda: FakeEmbedder())


def _add(conn, settings, collection, name, data, folder=None, user=None):
    received = store.receive(io.BytesIO(data), name, settings)
    return store.add_file(conn, settings, received, collection, user, folder_id=folder)[0]


def _bytes(conn, settings, file_id):
    return store.content_path(conn, settings, file_id).read_bytes()


# ------------------------------------------------------------ Agents folder
def test_every_inbox_gets_an_agents_folder(conn, settings, world):
    assert editing.ensure_all_agents(conn) == 2          # alice and bob
    agents = editing.agents_folder(conn, world["alice"], create=False)
    row = conn.execute("SELECT * FROM folders WHERE id = ?", (agents,)).fetchone()
    assert (row["name"], row["role"], row["collection_id"], row["parent_id"]) == ("Agents", "agents", world["inbox"], None)
    assert editing.ensure_all_agents(conn) == 0          # already there
    # Binned: a new one is made on next use.
    fs.trash(conn, fs.Access.of(conn, world["alice"]), [fs.Item("folder", agents)])
    again = editing.agents_folder(conn, world["alice"])
    assert again != agents and conn.execute("SELECT role FROM folders WHERE id = ?", (again,)).fetchone()[0] == "agents"


def test_an_existing_agents_folder_is_adopted(conn, settings, world):
    mine = fs.create_folder(conn, fs.Access.of(conn, world["alice"]), world["inbox"], None, "agents")
    assert editing.agents_folder(conn, world["alice"]) == mine


def test_agents_folder_is_shared_with_claude_even_when_the_inbox_is_not(conn, settings, world):
    agents = editing.agents_folder(conn, world["alice"])
    other = fs.create_folder(conn, fs.Access.of(conn, world["alice"]), world["inbox"], None, "Private")
    permissions.set_claude_enabled(conn, world["inbox"], False, world["alice"])
    claude = fs.Access.for_claude(conn, world["alice"])
    assert claude.can_read(world["inbox"], agents) and not claude.can_read(world["inbox"], other)
    assert not claude.can_read(world["inbox"], None)


# -------------------------------------------------------------- permissions
def test_who_may_edit_what(conn, settings, world):
    a, b = world["alice"], world["bob"]
    agents = editing.agents_folder(conn, a)
    sub = fs.create_folder(conn, fs.Access.of(conn, a), world["inbox"], agents, "Drafts")
    check = editing.check_edit
    assert check(conn, a, True, world["inbox"], agents).reason == "agents"
    assert check(conn, a, True, world["inbox"], sub).reason == "agents"          # subfolders too
    with pytest.raises(editing.EditForbidden, match="can only search"):
        check(conn, a, False, world["inbox"], agents)                           # search-only token
    with pytest.raises(editing.EditForbidden, match="outside the Agents folder"):
        check(conn, a, True, world["inbox"], None)
    with pytest.raises(editing.EditForbidden, match="this library"):
        check(conn, a, True, world["brand"], None)                              # library switch is off
    conn.execute("UPDATE collections SET claude_edit = 1 WHERE id = ?", (world["brand"],))
    assert check(conn, a, True, world["brand"], None).reason == "library"
    with pytest.raises(editing.EditForbidden, match="can't change files here yourself"):
        check(conn, b, True, world["brand"], None)                              # bob may only read
    with pytest.raises(editing.EditNotFound):
        check(conn, b, True, world["inbox"], agents)                            # alice's inbox
    permissions.set_claude_enabled(conn, world["brand"], False, a)
    with pytest.raises(editing.EditNotFound):
        check(conn, a, True, world["brand"], None)                              # not shared with Claude


def test_claude_makes_folders_where_it_may_edit(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    made, _ = tools.create_folder(conn, alice, "Red Lion")
    assert (made["library"], made["folder"], made["created"]) == ("File box", "Agents/Red Lion", True)
    assert made["link"] == f"/files#c={world['inbox']}&f={made['folder_id']}"
    again, _ = tools.create_folder(conn, alice, "red lion")                     # asking twice is harmless
    assert again["folder_id"] == made["folder_id"] and not again["created"] and "already" in again["note"]
    # create_file takes the folder it returned.
    out, _ = tools.create_file(conn, alice, "note.md", content="hi", library=made["library"], folder=made["folder"])
    assert out["where"] == "File box / Agents/Red Lion"

    with pytest.raises(editing.EditForbidden, match="this library"):
        tools.create_folder(conn, alice, "Agents", library="Brand")              # library switch is off
    conn.execute("UPDATE collections SET claude_edit = 1 WHERE id = ?", (world["brand"],))
    top, _ = tools.create_folder(conn, alice, "Agents", library="Brand")
    assert (top["library"], top["folder"]) == ("Brand", "Agents")
    with pytest.raises(editing.EditForbidden, match="outside the Agents folder"):
        tools.create_folder(conn, alice, "Loose", library="File box")            # File box root
    with pytest.raises(editing.EditForbidden, match="can only search"):
        tools.create_folder(conn, _caller(conn, world["alice"], can_edit=False), "Nope")
    with pytest.raises(editing.EditError):
        tools.create_folder(conn, alice, "bad/name")


# ---------------------------------------------------------------- text flow
def test_create_open_edit_conflict_versions_restore(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    made, _ = tools.create_file(conn, alice, "plan.md", content="# Plan\n\nCheck-in at 3 pm.\n")
    row = fs.file_row(conn, made["file_id"])
    assert made["where"] == "File box / Agents" and row["folder_id"] == editing.agents_folder(conn, world["alice"])
    assert row["edited_via"] == "claude:alice"

    opened, _ = tools.open_file(conn, alice, made["file_id"])
    assert opened["can_edit"] and opened["format"] == "text" and "3| Check-in at 3 pm." in opened["text"]
    edited, _ = tools.edit_text(conn, alice, made["file_id"], opened["version"],
                                edits=[{"old": "3 pm", "new": "4 pm"}])
    assert _bytes(conn, settings, made["file_id"]) == b"# Plan\n\nCheck-in at 4 pm.\n"
    assert edited["changes"] == ["replaced '3 pm' with '4 pm'"]
    assert fs.file_row(conn, made["file_id"])["status"] == "queued"            # re-indexed
    assert conn.execute("SELECT count(*) FROM jobs WHERE file_id = ?", (made["file_id"],)).fetchone()[0] >= 1

    with pytest.raises(editing.EditConflict, match="changed since"):
        tools.edit_text(conn, alice, made["file_id"], opened["version"], edits=[{"old": "Plan", "new": "X"}])
    with pytest.raises(editing.EditError, match="appears 0 times"):
        tools.edit_text(conn, alice, made["file_id"], edited["version"], edits=[{"old": "nope", "new": "x"}])

    versions = tools.list_versions(conn, alice, made["file_id"])[0]["earlier_versions"]
    assert len(versions) == 1 and versions[0]["replaced_by"] == "Claude for alice"
    assert "3 pm" in versions[0]["change"]
    editing.restore(conn, settings, made["file_id"], versions[0]["id"], world["alice"])
    assert _bytes(conn, settings, made["file_id"]) == b"# Plan\n\nCheck-in at 3 pm.\n"
    assert len(editing.versions(conn, made["file_id"])) == 2                  # the restore can be undone too


def test_text_keeps_bom_and_windows_line_endings():
    data = b"\xef\xbb\xbfone\r\ntwo\r\nthree\r\n"
    out = editing_formats.text_apply(data, [{"old": "two", "new": "2"}]).data
    assert out == b"\xef\xbb\xbfone\r\n2\r\nthree\r\n"
    view = editing_formats.text_view(data, start_line=2, max_lines=1)
    assert view["line_endings"] == "windows" and view["text"] == "2| two"
    with pytest.raises(editing.EditError, match="UTF-8"):
        editing_formats.text_apply(b"\xff\xfe\x00bad", [{"old": "a", "new": "b"}])


def test_search_only_tokens_and_wrong_files_are_refused(conn, settings, world, tools):
    viewer = _caller(conn, world["alice"], can_edit=False)
    with pytest.raises(editing.EditForbidden, match="can only search"):
        tools.create_file(conn, viewer, "x.txt", content="x")
    pdf = _add(conn, settings, world["inbox"], "a.pdf", b"%PDF-1.4 not really")
    with pytest.raises(editing.EditError, match="isn't a file Claude can open"):
        tools.open_file(conn, viewer, pdf)
    with pytest.raises(editing.EditError, match=r"\.txt, \.md, \.docx or \.xlsx"):
        tools.create_file(conn, _caller(conn, world["alice"]), "macro.docm", content="x")
    with pytest.raises(editing.EditError, match="Claude can create"):
        editing.create(conn, settings, _caller(conn, world["alice"]), "macro.xlsm", b"x")
    opened = tools.open_file(conn, viewer, _add(conn, settings, world["inbox"], "n.txt", b"hi"))[0]
    assert not opened["can_edit"] and "can only search" in opened["why_not"]


def test_library_editing_needs_the_switch_and_write_access(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    with pytest.raises(editing.EditForbidden):
        tools.create_file(conn, alice, "note.txt", content="x", library="Brand")
    conn.execute("UPDATE collections SET claude_edit = 1 WHERE id = ?", (world["brand"],))
    made = tools.create_file(conn, alice, "note.txt", content="x", library="Brand")[0]
    assert made["where"] == "Brand"
    with pytest.raises(editing.EditForbidden):
        tools.create_file(conn, _caller(conn, world["bob"]), "bob.txt", content="x", library="Brand")
    assert tools.list_libraries(conn, alice)[0]["you_can_edit_in"] == ["File box / Agents", "Brand"]


# --------------------------------------------------------------------- docx
def _docx_bytes():
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Eternal Hotels confidential"
    d.add_heading("Check-in", level=1)
    p = d.add_paragraph("Guests may check in from ")
    bold = p.add_run("3 pm")
    bold.bold = True
    p.add_run(" every day.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Room", "Rate"
    table.cell(1, 0).text, table.cell(1, 1).text = "King", "129"
    d.add_paragraph("Last line.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_docx_view_and_edits_keep_formatting(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    agents = editing.agents_folder(conn, world["alice"])
    file_id = _add(conn, settings, world["inbox"], "policy.docx", _docx_bytes(), folder=agents)
    view = tools.open_file(conn, alice, file_id)[0]
    assert [b["id"] for b in view["blocks"]] == ["p1", "p2", "t1", "p3"]
    assert view["blocks"][0] == {"id": "p1", "style": "Heading 1", "text": "Check-in"}
    assert view["blocks"][2]["table"][1][1] == {"id": "t1.r2.c2", "text": "129"}

    result = tools.edit_document(conn, alice, file_id, view["version"], [
        {"op": "replace_text", "paragraph": "p2", "old": "3 pm", "new": "4 pm"},
        {"op": "insert_paragraph", "after": "p2", "text": "Early check-in costs extra."},
        {"op": "set_cell", "cell": "t1.r2.c2", "text": "139"},
        {"op": "add_table_row", "table": "t1", "cells": ["Queen", "119"]},
        {"op": "delete_paragraph", "paragraph": "p3"},
        {"op": "insert_paragraph", "after": "end", "text": "Updated by Claude.", "style": "Heading 2"},
    ])[0]
    assert len(result["changes"]) == 6
    d = docx.Document(io.BytesIO(_bytes(conn, settings, file_id)))
    texts = [p.text for p in d.paragraphs]
    assert texts == ["Check-in", "Guests may check in from 4 pm every day.", "Early check-in costs extra.",
                     "Updated by Claude."]
    runs = d.paragraphs[1].runs
    assert [r.text for r in runs] == ["Guests may check in from ", "4 pm", " every day."] and runs[1].bold
    assert d.paragraphs[2].style.name == d.paragraphs[1].style.name     # inserted text matches its neighbour
    assert d.paragraphs[3].style.name == "Heading 2"
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == [["Room", "Rate"], ["King", "139"],
                                                                         ["Queen", "119"]]
    assert d.sections[0].header.paragraphs[0].text == "Eternal Hotels confidential"   # untouched parts stay

    with pytest.raises(editing.EditError, match="no paragraph 'p9'"):
        editing_formats.docx_apply(_docx_bytes(), [{"op": "delete_paragraph", "paragraph": "p9"}])
    with pytest.raises(editing.EditError, match="no style"):
        editing_formats.docx_apply(_docx_bytes(), [{"op": "set_paragraph", "paragraph": "p1", "text": "x",
                                                    "style": "Nonsense"}])
    with pytest.raises(editing.EditError, match="appears 0 times"):
        editing_formats.docx_apply(_docx_bytes(), [{"op": "replace_text", "paragraph": "p2", "old": "9 pm"}])
    with pytest.raises(editing.EditError, match="At most"):
        editing_formats.docx_apply(_docx_bytes(), [{"op": "delete_paragraph", "paragraph": "p1"}] * 51)


def test_new_docx_from_paragraphs(conn, settings, world, tools):
    made = tools.create_file(conn, _caller(conn, world["alice"]), "memo.docx",
                             paragraphs=[{"text": "Memo", "style": "Title"}, "Body text."])[0]
    d = docx.Document(io.BytesIO(_bytes(conn, settings, made["file_id"])))
    assert [(p.text, p.style.name) for p in d.paragraphs] == [("Memo", "Title"), ("Body text.", "Normal")]


def test_any_kind_of_file_claude_made_is_saved_from_base64(conn, settings, world, tools):
    """Images, PDFs, decks: create_file's `data` (base64, a data: URL too) stores the bytes as they are."""
    import base64
    import PIL.Image
    from mnemerynne import agent as agent_mod, proposals, users

    alice = _caller(conn, world["alice"])
    png = io.BytesIO()
    PIL.Image.new("RGB", (4, 4), "pink").save(png, "PNG")
    made = tools.create_file(conn, alice, "chart.png", data=base64.b64encode(png.getvalue()).decode())[0]
    assert _bytes(conn, settings, made["file_id"]) == png.getvalue()
    assert made["where"] == "File box / Agents" and made["kind"] == "other" and made["size"] == len(png.getvalue())
    assert made["version"] == editing.sha256(png.getvalue())
    row = conn.execute("SELECT mime, edited_via FROM files f JOIN blobs b ON b.sha256 = f.blob_sha256 WHERE f.id = ?",
                       (made["file_id"],)).fetchone()
    assert row["mime"] == "image/png" and row["edited_via"] == "claude:alice"
    # A data: URL, whitespace in the base64, and a name without a known editable extension.
    pdf = b"%PDF-1.4 not really"
    url = "data:application/pdf;base64," + base64.b64encode(pdf).decode()
    url = url[:20] + "\n" + url[20:]
    made = tools.create_file(conn, alice, "deck.pdf", data=url)[0]
    assert _bytes(conn, settings, made["file_id"]) == pdf and made["kind"] in ("pdf", "other")
    # The same upload rules as a browser: size, emptiness, and the usual places.
    with pytest.raises(editing.EditError, match="valid base64"):
        tools.create_file(conn, alice, "x.bin", data="not*base64")
    with pytest.raises(editing.EditError, match="empty"):
        tools.create_file(conn, alice, "x.bin", data="")
    with pytest.raises(editing.EditError, match="not both"):
        tools.create_file(conn, alice, "x.md", data="aGk=", content="hi")
    with pytest.raises(editing.EditError, match="`data`"):
        tools.create_file(conn, alice, "x.zip", content="hi")
    with pytest.raises(editing.EditError, match="MB"):
        settings.edit_max_bytes = 8
        tools.create_file(conn, alice, "big.bin", data=base64.b64encode(b"x" * 9).decode())
    settings.edit_max_bytes = 20 * 1024 * 1024
    with pytest.raises(editing.EditError):
        tools.create_file(conn, _caller(conn, world["alice"], can_edit=False), "x.bin", data="aGk=")
    # Erynne proposes it; the person applies, and the bytes land unchanged.
    erynne = Tools(settings, lambda: FakeEmbedder(), mode="erynne")
    caller = agent_mod.AgentCaller(world["alice"], "alice", None)
    proposed = erynne.create_file(conn, caller, "pic.png", data=base64.b64encode(png.getvalue()).decode())[0]
    assert proposed["status"] == "proposed"
    applied = proposals.apply(conn, settings, users.get_user_by_id(conn, world["alice"]), proposed["proposal"])
    assert _bytes(conn, settings, applied["file_id"]) == png.getvalue()


# --------------------------------------------------------------------- xlsx
def _xlsx_bytes(formulas=True, chart=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["Room", "Rate"])
    ws.append(["King", 129])
    ws.append(["Queen", 119])
    if formulas:
        ws["B4"] = "=SUM(B2:B3)"
    if chart:
        from openpyxl.chart import BarChart, Reference
        c = BarChart()
        c.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        ws.add_chart(c, "D2")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_xlsx_view_and_edits(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    agents = editing.agents_folder(conn, world["alice"])
    file_id = _add(conn, settings, world["inbox"], "rates.xlsx", _xlsx_bytes(), folder=agents)
    view = tools.open_file(conn, alice, file_id)[0]
    assert view["editable"] and view["sheet"] == "Rates" and view["sheets"][0]["rows"] == 4
    assert view["cells"][1] == {"A2": "King", "B2": 129}
    assert view["cells"][3]["B4"]["formula"] == "=SUM(B2:B3)"

    tools.edit_spreadsheet(conn, alice, file_id, view["version"], [
        {"op": "set_cells", "cells": {"B2": 139, "C1": "Notes", "C2": "=B2*1.1"}},
        {"op": "append_rows", "rows": [["Suite", 249]]},
        {"op": "add_sheet", "name": "Summary"},
    ])
    wb = openpyxl.load_workbook(io.BytesIO(_bytes(conn, settings, file_id)))
    ws = wb["Rates"]
    assert (ws["B2"].value, ws["C1"].value, ws["C2"].value, ws["A5"].value, ws["B5"].value) == \
        (139, "Notes", "=B2*1.1", "Suite", 249)
    assert wb.sheetnames == ["Rates", "Summary"] and wb.calculation.fullCalcOnLoad

    with pytest.raises(editing.EditError, match="has formulas"):
        editing_formats.xlsx_apply(_xlsx_bytes(), [{"op": "insert_rows", "row": 2}])
    plain = editing_formats.xlsx_apply(_xlsx_bytes(formulas=False), [{"op": "delete_rows", "row": 2}]).data
    assert openpyxl.load_workbook(io.BytesIO(plain))["Rates"]["A2"].value == "Queen"
    renamed = openpyxl.Workbook()
    renamed.active.title = "Rates"
    renamed.create_sheet("Totals")["A1"] = "=Rates!B2*2"
    out = io.BytesIO()
    renamed.save(out)
    with pytest.raises(editing.EditError, match="formulas refer to 'Rates'"):
        editing_formats.xlsx_apply(out.getvalue(), [{"op": "rename_sheet", "sheet": "Rates", "name": "Prices"}])


def test_workbooks_that_would_be_damaged_are_refused(conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    agents = editing.agents_folder(conn, world["alice"])
    file_id = _add(conn, settings, world["inbox"], "chart.xlsx", _xlsx_bytes(chart=True), folder=agents)
    view = tools.open_file(conn, alice, file_id)[0]
    assert not view["editable"] and "charts" in view["note"]
    before = _bytes(conn, settings, file_id)
    with pytest.raises(editing.EditError, match="charts"):
        tools.edit_spreadsheet(conn, alice, file_id, view["version"], [{"op": "set_cells", "cells": {"B2": 1}}])
    assert _bytes(conn, settings, file_id) == before


# ------------------------------------------------------------- housekeeping
def test_old_versions_are_dropped(conn, settings, world):
    settings.version_keep = 2
    agents = editing.agents_folder(conn, world["alice"])
    file_id = _add(conn, settings, world["inbox"], "log.txt", b"v0", folder=agents)
    for n in range(1, 5):
        editing.replace_content(conn, settings, file_id, f"v{n}".encode(), expected=None, actor_id=world["alice"],
                                via="claude", summary=f"v{n}")
    assert len(editing.versions(conn, file_id)) == 4
    assert editing.purge_versions(conn, settings) == 2
    kept = editing.versions(conn, file_id)
    assert [v["change"] for v in kept] == ["v4", "v3"]
    conn.execute("UPDATE file_versions SET created_at = datetime('now', '-40 days')")
    editing.purge_versions(conn, settings)
    assert editing.versions(conn, file_id) == []
    assert not (settings.data_dir / "versions" / str(file_id)).exists()


# ------------------------------------------------------------- on the share
@pytest.fixture
def share_world(tmp_path):
    root = tmp_path / "storage" / "share"
    root.mkdir(parents=True)
    (root / share.MARKER).write_text("")
    settings = Settings(_env_file=None, data_dir=tmp_path / "data", secret_key="test-secret-key-for-the-test-suite-only", embed_dim=4,
                        ad_domain="ETERNALHOTELS", share_root=root, share_settle_seconds=0)
    directory = FakeDirectory({"jsmith": {"Domain Users"}})
    share.configure(settings, directory, FakeShareChecker(root))
    connection = db.connect(settings.db_path)
    db.migrate(connection, settings.embed_dim)
    jane = users.create_user(connection, "jane", "password123")
    accounts.link(connection, jane, "jsmith", directory)
    inbox = permissions.private_collection_id(connection, jane)
    if not share.is_share(connection, inbox):
        share.attach_inbox(connection, inbox, "jane", "jsmith")
    yield {"conn": connection, "settings": settings, "root": root, "jane": jane, "inbox": inbox}
    connection.close()


def test_editing_on_the_share(share_world):
    conn, settings, root, jane = (share_world[k] for k in ("conn", "settings", "root", "jane"))
    tools = Tools(settings, lambda: FakeEmbedder())
    caller = _caller(conn, jane)
    made = tools.create_file(conn, caller, "notes.md", content="Pool opens at 8.")[0]
    on_disk = root / "Users" / "jane" / "Agents" / "notes.md"
    assert on_disk.read_bytes() == b"Pool opens at 8."

    opened = tools.open_file(conn, caller, made["file_id"])[0]
    tools.edit_text(conn, caller, made["file_id"], opened["version"], edits=[{"old": "8", "new": "7"}])
    assert on_disk.read_bytes() == b"Pool opens at 7."
    assert not [p for p in on_disk.parent.iterdir() if p.name.startswith(".mnemerynne-")]   # no temp left
    assert share.sync(conn, settings).changed == 0              # the index already knows the new content

    # Someone saves it over SMB meanwhile: the old version is refused.
    stale = tools.open_file(conn, caller, made["file_id"])[0]["version"]
    on_disk.write_bytes(b"Pool opens at 6.")
    with pytest.raises(editing.EditConflict, match="changed since"):
        tools.edit_text(conn, caller, made["file_id"], stale, edits=[{"old": "6", "new": "5"}])

    # Open in Word right now: refused.
    fresh = tools.open_file(conn, caller, made["file_id"])[0]["version"]
    lock = on_disk.with_name("~$notes.md")
    lock.write_bytes(b"lock")
    with pytest.raises(editing.EditConflict, match="open in Word or Excel"):
        tools.edit_text(conn, caller, made["file_id"], fresh, edits=[{"old": "6", "new": "5"}])
    lock.unlink()
    tools.edit_text(conn, caller, made["file_id"], fresh, edits=[{"old": "6", "new": "5"}])
    assert on_disk.read_bytes() == b"Pool opens at 5."
    assert [v["change"] for v in editing.versions(conn, made["file_id"])][0] == "replaced '6' with '5'"


# ------------------------------------------------------------------- web/API
@pytest.fixture
def web(settings, conn, world):
    users.create_user(conn, "boss", "password123", role="admin")
    with TestClient(create_app(settings, Services(llm=None, embedder=FakeEmbedder()))) as client:
        yield client


def _login(client, name):
    client.post("/login", data={"username": name, "password": "password123"})


def test_versions_and_restore_in_files(web, conn, settings, world, tools):
    alice = _caller(conn, world["alice"])
    made = tools.create_file(conn, alice, "a.txt", content="one")[0]
    v = tools.open_file(conn, alice, made["file_id"])[0]["version"]
    tools.edit_text(conn, alice, made["file_id"], v, content="two")
    _login(web, "alice")
    info = web.get(f"/api/files/{made['file_id']}").json()
    assert info["edited_via"] == "claude:alice"
    listed = web.get(f"/api/files/{made['file_id']}/versions").json()
    assert listed["can_restore"] and len(listed["versions"]) == 1
    old = listed["versions"][0]
    download = web.get(old["download_url"])
    assert download.content == b"one" and "a (before" in unquote(download.headers["content-disposition"])
    restored = web.post(f"/api/files/{made['file_id']}/versions/{old['id']}/restore")
    assert restored.status_code == 200 and restored.json()["edited_via"] is None
    assert _bytes(conn, settings, made["file_id"]) == b"one"
    web.post("/logout")
    _login(web, "bob")
    assert web.get(f"/api/files/{made['file_id']}/versions").status_code == 404
    assert web.post(f"/api/files/{made['file_id']}/versions/{old['id']}/restore").status_code == 404


def test_token_edit_setting_and_admin_switch(web, conn, world):
    _login(web, "alice")
    token = web.post("/api/tokens", json={"name": "Desktop", "days": 30}).json()
    caller = api_tokens.verify(conn, token["token"])
    assert caller.can_edit                                         # new tokens may edit by default
    assert web.patch(f"/api/tokens/{token['id']}", json={"can_edit": False}).json() == {"id": token["id"],
                                                                                         "can_edit": False}
    assert not api_tokens.verify(conn, token["token"]).can_edit
    listed = web.get("/api/tokens").json()
    assert listed["tokens"][0]["can_edit"] is False
    assert listed["editable"] == [{"library": "File box", "folder": "Agents", "why": "your Agents folder"}]
    assert web.put(f"/api/admin/libraries/{world['brand']}/claude-edit", json={"enabled": True}).status_code == 403
    web.post("/logout")
    _login(web, "boss")
    assert web.put(f"/api/admin/libraries/{world['brand']}/claude-edit", json={"enabled": True}).json()["claude_edit"]
    libs = {l["name"]: l for l in web.get("/api/admin/libraries").json()["libraries"]}
    assert libs["Brand"]["claude_edit"] is True
    assert web.put(f"/api/admin/users/{world['alice']}/inbox-claude-edit", json={"enabled": True}).status_code == 200
    assert web.get("/api/admin/log").json()["entries"][0]["action"] == "claude.edit.on"
    web.post("/logout")
    _login(web, "alice")
    assert {(p["library"], p["why"]) for p in web.get("/api/tokens").json()["editable"]} >= {("Brand", "allowed by an admin")}


def test_edit_tools_over_mcp_are_logged(web, conn, world):
    _, token = api_tokens.create(conn, world["alice"], "t", can_edit=True)
    headers = {"Accept": "application/json, text/event-stream", "Authorization": f"Bearer {token}"}
    call = lambda name, args: web.post("/mcp", headers=headers, json={  # noqa: E731
        "jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": name, "arguments": args}}).json()["result"]
    made = call("create_file", {"name": "todo.md", "content": "- [ ] audit"})
    assert not made.get("isError"), made
    payload = made.get("structuredContent") or json.loads(made["content"][0]["text"])
    refused = call("edit_text", {"file_id": payload["file_id"], "version": "stale",
                                 "edits": [{"old": "audit", "new": "night audit"}]})
    assert refused["isError"] and "changed since" in refused["content"][0]["text"]
    logged = conn.execute("SELECT tool, ok FROM mcp_calls ORDER BY id").fetchall()
    assert [(r["tool"], r["ok"]) for r in logged] == [("create_file", 1), ("edit_text", 0)]
